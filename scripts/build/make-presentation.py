# © 2026 Dea Saka Kurnia Putra. All rights reserved.
# Builds docs/presentation/Finandex-Presentation-{ID,EN}.pptx (16:9, Finandex brand theme) with python-pptx.
# Screens come from scripts/build/capture-screens.ts. Run: python3 -I scripts/build/make-presentation.py
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

ROOT = os.getcwd()
BG = RGBColor(0x0A, 0x0F, 0x0C)
SURF = RGBColor(0x11, 0x19, 0x15)
SURF2 = RGBColor(0x17, 0x22, 0x1C)
TEXT = RGBColor(0xEE, 0xF3, 0xEF)
MUTED = RGBColor(0x9D, 0xB0, 0xA4)
GREEN = RGBColor(0x19, 0xF2, 0x3A)
AMBER = RGBColor(0xFF, 0xB0, 0x20)
BLUE = RGBColor(0x4D, 0xA3, 0xFF)
HEAD = 'Arial'
BODY = 'Arial'
W, H = Inches(13.333), Inches(7.5)

DEX = [
    ('Pasar Global & Saham', 'Global Markets & Equities', ['Finandex AI', 'NASDAQDex', 'SnP500Dex', 'CommerceDex', 'IPODex']),
    ('Pasar & Ekonomi Indonesia', 'Indonesian Markets & Economy', ['IDFinanDex', 'RupiahDex', 'IDXDex', 'ReksaDex']),
    ('Kripto & Makro Global', 'Crypto & Global Macro', ['CryptDex', 'CentralBankDex', 'InflationDex']),
    ('Startup, VC & Perusahaan', 'Startups, VC & Companies', ['StartupDex', 'VCDex', 'VCNetworkDex', 'OpenCompanyDex']),
    ('Biaya Hidup, Kerja & Rumah Tangga', 'Cost of Living, Work & Household', ['SalaryDex', 'CostLivingDex', 'IDHousingDex', 'IDFoodDeliveryDex']),
    ('Sosial & Pembangunan', 'Social & Development', ['PovertyDex', 'PopulationDex', 'HomelessnessDex', 'HousingDex']),
    ('Sentimen & Keuangan Pribadi', 'Sentiment & Personal Finance', ['FintechDex', 'PersonalFinanDex']),
    ('Hukum', 'Law', ['My LAIwyer']),
]
LAW = json.load(open(os.path.join(ROOT, 'src/config/law-finance-sources.json')))['docs']


def bg(slide):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = BG


def box(slide, x, y, w, h, fill=SURF, line=None, radius=True):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    if radius:
        s.adjustments[0] = 0.08
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, runs, size=16, color=TEXT, font=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = runs if isinstance(runs, list) else [runs]
    for i, ln in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        par.line_spacing = spacing
        r = par.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = font
        r.font.bold = bold
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=TEXT, gap=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.space_after = Pt(gap)
        par.line_spacing = 1.12
        r1 = par.add_run()
        r1.text = '■  '
        r1.font.size = Pt(size - 6)
        r1.font.color.rgb = GREEN
        r2 = par.add_run()
        r2.text = it
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
        r2.font.name = BODY
    return tb


def title(slide, kicker, heading, sub=None):
    bg(slide)
    text(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.4), kicker.upper(), 12, GREEN, BODY, True)
    text(slide, Inches(0.6), Inches(0.72), Inches(12.1), Inches(0.9), heading, 32, TEXT, HEAD, True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.55), Inches(1.2), Emu(38000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()
    if sub:
        text(slide, Inches(0.6), Inches(1.7), Inches(12), Inches(0.6), sub, 15, MUTED)


def footer(slide, n, L):
    text(slide, Inches(0.6), Inches(7.05), Inches(9), Inches(0.3), '© 2026 Dea Saka Kurnia Putra · Finandex', 10, MUTED)
    text(slide, Inches(11.8), Inches(7.05), Inches(1), Inches(0.3), str(n), 10, MUTED, align=PP_ALIGN.RIGHT)


def stat(slide, x, y, value, label, color=GREEN):
    box(slide, x, y, Inches(2.3), Inches(1.6), SURF2)
    text(slide, x + Inches(0.15), y + Inches(0.18), Inches(2.0), Inches(0.8), value, 34, color, HEAD, True)
    text(slide, x + Inches(0.15), y + Inches(0.95), Inches(2.0), Inches(0.6), label, 12, MUTED)


def table(slide, x, y, w, rows, col_w, size=11, head_fill=SURF2, row_h=Inches(0.38)):
    t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows)).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = cw
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = head_fill if i == 0 else (SURF if i % 2 else BG)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ''
            r_ = p.add_run()
            r_.text = v
            r_.font.size = Pt(size)
            r_.font.name = BODY
            r_.font.bold = i == 0
            r_.font.color.rgb = GREEN if i == 0 else TEXT
    return t


def screen(slide, path, x, y, w):
    box(slide, x - Inches(0.06), y - Inches(0.06), w + Inches(0.12), int(w * 900 / 1440) + Inches(0.12), SURF2, GREEN)
    slide.shapes.add_picture(path, x, y, width=w)


def flow(slide, y, steps, w=Inches(2.25), color=GREEN):
    x = Inches(0.6)
    for i, (head, body) in enumerate(steps):
        box(slide, x, y, w, Inches(1.9), SURF2, color)
        text(slide, x + Inches(0.12), y + Inches(0.12), w - Inches(0.24), Inches(0.5), head, 15, color, HEAD, True)
        text(slide, x + Inches(0.12), y + Inches(0.6), w - Inches(0.24), Inches(1.25), body, 11.5, TEXT)
        if i < len(steps) - 1:
            a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + w + Inches(0.03), y + Inches(0.8), Inches(0.24), Inches(0.3))
            a.fill.solid()
            a.fill.fore_color.rgb = MUTED
            a.line.fill.background()
        x += w + Inches(0.3)


def build(L):
    I = L == 'id'
    t = lambda a, b: a if I else b
    sc = lambda n: os.path.join(ROOT, f'docs/presentation/screens/{n}-{L}.png')
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    n = 0

    def new():
        nonlocal n
        n += 1
        s = prs.slides.add_slide(blank)
        return s

    # 1 Cover
    s = new()
    bg(s)
    s.shapes.add_picture(os.path.join(ROOT, 'kit/brand/icons/finandex-1024.png'), Inches(0.8), Inches(0.9), width=Inches(1.3))
    text(s, Inches(0.8), Inches(2.5), Inches(11), Inches(1.2), 'Finandex', 60, TEXT, HEAD, True)
    text(s, Inches(0.8), Inches(3.55), Inches(11), Inches(0.7), t('Wawasan Indeks Pasar Keuangan Global', 'Global Financial Market Indexing Insights'), 24, GREEN, HEAD, True)
    text(s, Inches(0.8), Inches(4.3), Inches(11), Inches(1.0), t('27 chatbot berbasis data untuk pasar, ekonomi, dan hukum — termasuk hukum keuangan dan bantuan korban penipuan', '27 data-grounded chatbots for markets, economics, and law — including financial law and help for scam victims'), 18, MUTED)
    s.shapes.add_picture(os.path.join(ROOT, 'kit/brand/dskp-mark-on-dark-192.png'), Inches(0.8), Inches(6.15), height=Inches(0.45))
    text(s, Inches(1.35), Inches(6.18), Inches(6), Inches(0.4), 'by Dea Saka Kurnia Putra', 14, TEXT)
    text(s, Inches(8.3), Inches(6.18), Inches(4.4), Inches(0.4), t('Oktober 2026 · finandex-beta.vercel.app', 'October 2026 · finandex-beta.vercel.app'), 12, MUTED, align=PP_ALIGN.RIGHT)

    # 2 Problem
    s = new()
    title(s, t('Masalah', 'Problem'), t('Data ada di mana-mana, kepercayaan tidak', 'Data is everywhere, trust is not'))
    cards = [
        (t('Tersebar', 'Scattered'), t('Harga, makro, startup, dan sosial tersebar di banyak sumber dengan format, kualitas, dan tanggal berbeda.', 'Prices, macro, startup, and social data live in many sources with different formats, quality, and dates.')),
        (t('Tidak dapat ditelusuri', 'Not traceable'), t('Chatbot umum lancar menjawab, tetapi jarang menunjukkan asal angka, tanggal data, atau cara menghitungnya.', 'General chatbots answer fluently but rarely show where a number came from, its date, or how it was computed.')),
        (t('Korban penipuan', 'Scam victims'), t('Penipuan transaksi keuangan terus naik; korban butuh langkah segera, kanal resmi, dan dasar hukum yang jelas.', 'Financial fraud keeps rising; victims need immediate steps, official channels, and a clear legal basis.')),
    ]
    for i, (h, b) in enumerate(cards):
        x = Inches(0.6 + i * 4.1)
        box(s, x, Inches(2.4), Inches(3.85), Inches(3.6), SURF2)
        text(s, x + Inches(0.25), Inches(2.6), Inches(3.4), Inches(0.6), h, 22, AMBER if i == 2 else GREEN, HEAD, True)
        text(s, x + Inches(0.25), Inches(3.35), Inches(3.4), Inches(2.5), b, 16, TEXT)
    footer(s, n, L)

    # 3 Solution
    s = new()
    title(s, t('Solusi', 'Solution'), t('Setiap jawaban dapat ditelusuri ke datanya', 'Every answer traces back to its data'))
    bullets(s, Inches(0.6), Inches(2.1), Inches(6.2), Inches(4.7), [
        t('Angka hanya dari dataset terverifikasi SHA-256', 'Numbers only from SHA-256-verified datasets'),
        t('Grafik, sumber, tanggal data, dan rumus di setiap jawaban', 'Chart, source, as-of date, and formula in every answer'),
        t('Satu jalur hitung: chat = tab fitur (tool yang sama)', 'One computation path: chat = feature tabs (same tools)'),
        t('Kutipan hukum dari teks resmi peraturan', 'Legal quotes from official regulation texts'),
        t('Tanpa akun, tanpa pelacakan, offline, dua bahasa', 'No accounts, no tracking, offline, two languages'),
    ], 18)
    screen(s, sc('chat'), Inches(7.1), Inches(2.15), Inches(5.6))
    footer(s, n, L)

    # 4 Key numbers
    s = new()
    title(s, t('Sekilas', 'At a glance'), t('Platform dalam angka', 'The platform in numbers'))
    stats = [('27', 'Dex'), ('42', t('dataset terverifikasi', 'verified datasets')), ('19', t('teks resmi hukum keuangan', 'official financial-law texts')), ('362', t('tool analitik read-only', 'read-only analytics tools')), ('793', t('test otomatis lulus', 'automated tests passing')), ('128/128', t('evaluasi chatbot (mock)', 'chatbot evaluation (mock)')), ('59/59', t('parity Python ↔ TypeScript', 'Python ↔ TypeScript parity')), ('180', t('panel fitur interaktif', 'interactive feature panels'))]
    for i, (v, l) in enumerate(stats):
        stat(s, Inches(0.6 + (i % 4) * 3.1), Inches(2.2 + (i // 4) * 2.1), v, l)
    footer(s, n, L)

    # 5 Dex map
    s = new()
    title(s, t('Produk', 'Product'), t('27 Dex dalam 8 kategori', '27 Dex in 8 categories'))
    for i, (gi, ge, items) in enumerate(DEX):
        x = Inches(0.6 + (i % 4) * 3.1)
        y = Inches(2.05 + (i // 4) * 2.45)
        box(s, x, y, Inches(2.95), Inches(2.25), SURF2)
        text(s, x + Inches(0.15), y + Inches(0.12), Inches(2.7), Inches(0.6), t(gi, ge), 13, GREEN, HEAD, True)
        text(s, x + Inches(0.15), y + Inches(0.75), Inches(2.7), Inches(1.5), items, 12.5, TEXT, spacing=1.05)
    footer(s, n, L)

    # 6 Home screen
    s = new()
    title(s, t('Pengalaman pengguna', 'User experience'), t('Beranda: peta seluruh fitur', 'Home: a map of every feature'))
    screen(s, sc('home'), Inches(0.6), Inches(2.0), Inches(7.6))
    bullets(s, Inches(8.6), Inches(2.0), Inches(4.2), Inches(4.8), [
        t('Kolom "Tanya Finandex AI" + saran lintas kategori', '"Ask Finandex AI" box + cross-category suggestions'),
        t('Ringkasan 6 aset dengan sparkline', '6-asset summary with sparklines'),
        t('Whitepaper: unduh PDF ID / EN', 'Whitepaper: download PDF ID / EN'),
        t('Peta fitur: cari, kategori, filter Indonesia', 'Feature map: search, categories, Indonesia filter'),
        t('Strip "Data per" semua Dex', '"As of" strip for every Dex'),
    ], 15)
    footer(s, n, L)

    # 7 Dex page + chat
    s = new()
    title(s, t('Halaman Dex & chat', 'Dex page & chat'), t('Tab fitur dan kartu jawaban yang menjelaskan dirinya', 'Feature tabs and self-explaining answer cards'))
    screen(s, sc('dex'), Inches(0.6), Inches(2.1), Inches(6.0))
    screen(s, sc('chat'), Inches(6.9), Inches(2.1), Inches(6.0))
    text(s, Inches(0.6), Inches(6.1), Inches(6), Inches(0.6), t('Tab fitur menjalankan tool yang sama dengan chatbot, langsung di browser.', 'Feature tabs run the same tools as the chatbot, right in the browser.'), 12.5, MUTED)
    text(s, Inches(6.9), Inches(6.1), Inches(6), Inches(0.6), t('Panel "Kenapa jawabannya begini?": data, periode, baris, rumus.', '"Why this answer?" panel: data, period, rows, formulas.'), 12.5, MUTED)
    footer(s, n, L)

    # 8 My LAIwyer scam help
    s = new()
    title(s, 'My LAIwyer', t('Bantuan korban penipuan keuangan', 'Help for financial-scam victims'))
    screen(s, sc('scam'), Inches(0.6), Inches(2.0), Inches(7.0))
    bullets(s, Inches(8.0), Inches(2.0), Inches(4.8), Inches(3.0), [
        t('8 skenario: OTP/phishing, investasi bodong, pinjol ilegal, salah transfer, transfer ke penipu, skimming, rekening penampung, kebocoran data', '8 scenarios: OTP/phishing, fraudulent investment, illegal loans, mistaken transfer, tricked transfer, skimming, money mule, data leak'),
        t('Langkah segera + kanal resmi', 'Immediate steps + official channels'),
        t('Dasar hukum dikutip dari teks peraturan', 'Legal basis quoted from the regulations'),
    ], 14)
    box(s, Inches(8.0), Inches(5.05), Inches(4.8), Inches(1.75), SURF2, AMBER)
    text(s, Inches(8.15), Inches(5.12), Inches(4.5), Inches(1.7), [t('Bank Anda → blokir', 'Your bank → block'), 'IASC · iasc.ojk.go.id', t('Kontak OJK 157 · Satgas PASTI', 'OJK Contact 157 · PASTI Task Force'), t('Polisi · BI Bicara 131 · Darurat 112', 'Police · BI Bicara 131 · Emergency 112')], 13, TEXT, spacing=1.1)
    footer(s, n, L)

    # 9 Law corpus
    s = new()
    title(s, 'My LAIwyer', t('Korpus hukum sektor keuangan: 19 teks resmi', 'Financial-sector law corpus: 19 official texts'))
    half = (len(LAW) + 1) // 2
    for k, chunk in enumerate([LAW[:half], LAW[half:]]):
        rows = [[t('Dokumen', 'Document'), t('Hal.', 'Pp.')]] + [[d['short'][L], str(d['pages'])] for d in chunk]
        table(s, Inches(0.6 + k * 6.25), Inches(2.0), Inches(6.0), rows, [Inches(5.2), Inches(0.8)], 11, row_h=Inches(0.42))
    text(s, Inches(0.6), Inches(6.65), Inches(12), Inches(0.4), t('Sumber: peraturan.go.id · ojk.go.id · JDIH BPK — SHA-256 dipatok; teks peraturan tidak dilindungi hak cipta (UU 28/2014 Pasal 42).', 'Sources: peraturan.go.id · ojk.go.id · JDIH BPK — pinned SHA-256; statutory texts are not copyrighted (Law 28/2014 Art. 42).'), 11, MUTED)
    footer(s, n, L)

    # 10 Data pipeline
    s = new()
    title(s, t('Data', 'Data'), t('Pipeline data terverifikasi', 'A verified data pipeline'))
    flow(s, Inches(2.3), [
        (t('Verifikasi', 'Verify'), t('42 dataset + 19 PDF hukum: ukuran, SHA-256, isi arsip.', '42 datasets + 19 legal PDFs: size, SHA-256, archive contents.')),
        (t('Ekstrak', 'Extract'), t('ZIP/CSV/Parquet/PDF streaming; konten pihak ketiga = data.', 'Streaming ZIP/CSV/Parquet/PDF; third-party content = data.')),
        (t('Bersihkan', 'Clean'), t('Spike, bad print FX, carry-forward, dedupe, privasi.', 'Spikes, FX bad prints, carry-forward, dedupe, privacy.')),
        (t('Validasi', 'Validate'), t('Skema zod, ≤ 10 digit, tanpa NaN; golden 62/62.', 'Zod schemas, ≤ 10 digits, no NaN; golden 62/62.')),
        (t('Sajikan', 'Serve'), t('core (offline) + heavy; manifest SHA-256.', 'core (offline) + heavy; SHA-256 manifest.')),
    ])
    text(s, Inches(0.6), Inches(4.6), Inches(12), Inches(1.5), t('Analitik diporting 1:1 dari referensi Python dan diuji terhadap nilai acuan (toleransi relatif 1e-6). Rumus yang sama tampil di panel penjelasan.', 'Analytics are ported 1:1 from the Python reference and tested against golden values (relative tolerance 1e-6). The same formulas appear in the explanation panel.'), 16, MUTED)
    footer(s, n, L)

    # 11 Architecture
    s = new()
    title(s, t('Arsitektur', 'Architecture'), t('Satu codebase, satu jalur angka', 'One codebase, one number path'))
    flow(s, Inches(2.2), [
        ('ETL', t('TypeScript, tsx: data/inputs → public/data', 'TypeScript, tsx: data/inputs → public/data')),
        (t('Data statis', 'Static data'), t('CDN + CORS; service worker cache', 'CDN + CORS; service-worker cache')),
        ('Tools', t('362 tool read-only, skema zod strict', '362 read-only tools, strict zod schemas')),
        (t('Server', 'Server'), t('/api/chat: validasi, rate limit, orkestrator, SSE', '/api/chat: validation, rate limit, orchestrator, SSE')),
        (t('Klien', 'Client'), t('React PWA + Capacitor; tab fitur jalankan tool di browser', 'React PWA + Capacitor; feature tabs run tools in the browser')),
    ], color=BLUE)
    table(s, Inches(0.6), Inches(4.5), Inches(12.1), [
        [t('Mode', 'Mode'), t('Cara kerja', 'How it works')],
        [t('AI (Claude)', 'AI (Claude)'), t('Model memilih tool dengan skema ketat dan menulis analisis hanya dari hasil tool (≤ 8 giliran, ≤ 12 tool).', 'The model picks tools with strict schemas and writes analysis only from tool results (≤ 8 turns, ≤ 12 tools).')],
        [t('Data (tanpa model)', 'Data (no model)'), t('Router deterministik: entitas, periode, peringkat, opsi, negara, skenario penipuan, topik hukum; jawaban disusun dari hasil tool.', 'Deterministic router: entities, periods, rankings, options, countries, scam scenarios, legal topics; answers assembled from tool results.')],
    ], [Inches(2.4), Inches(9.7)], 12.5, row_h=Inches(0.55))
    footer(s, n, L)

    # 12 Tech stack
    s = new()
    title(s, 'Tech stack', t('Teknologi yang dipakai', 'Technology used'))
    rows = [[t('Lapisan', 'Layer'), t('Teknologi', 'Technology')],
            [t('Bahasa & runtime', 'Language & runtime'), 'TypeScript 6.0 (strict) · Node.js 24 LTS'],
            [t('Antarmuka', 'Interface'), 'React 19.3 · React Router 8.5 · Vite 8.3 · Tailwind CSS 4.3 · Radix UI'],
            [t('Grafik & dokumen', 'Charts & documents'), 'Apache ECharts 6.1 · pdf.js 6.4 (WASM CCITT/JBIG2)'],
            [t('Data klien', 'Client data'), 'TanStack Query 5 · Zustand 5 · IndexedDB (idb-keyval) · TanStack Virtual'],
            [t('Skema & pencarian', 'Schemas & search'), 'Zod 4.6 · MiniSearch 7.2 (BM25)'],
            ['ETL', 'tsx · PapaParse · hyparquet · fflate · pdf.js'],
            [t('Server & AI', 'Server & AI'), 'Vercel/Netlify Functions (Web Fetch API) · esbuild · Anthropic Claude SDK 1.0'],
            [t('PWA & native', 'PWA & native'), 'vite-plugin-pwa 2.0 · Workbox 7.4 · Capacitor 8.5 (Android, iOS)'],
            [t('Kualitas', 'Quality'), 'Vitest 5 · Playwright 1.64 · axe-core · ESLint 9 · size-limit · Lighthouse CI'],
            [t('Infrastruktur', 'Infrastructure'), 'GitHub (repo, Actions, Releases) · Vercel · Netlify']]
    table(s, Inches(0.6), Inches(2.0), Inches(12.1), rows, [Inches(3.0), Inches(9.1)], 13, row_h=Inches(0.43))
    footer(s, n, L)

    # 13 Quality
    s = new()
    title(s, t('Kualitas', 'Quality'), t('Gate kualitas "tanpa cacat"', '"Zero-defect" quality gates'))
    q = [('0', t('error lint & typecheck', 'lint & typecheck errors')), ('793', t('test lulus · cakupan 96%', 'tests passing · 96% coverage')), ('62/62', t('golden test', 'golden tests')), ('128/128', t('eval chatbot · keamanan 100%', 'chatbot eval · safety 100%')),
         ('324', t('saran pertanyaan dijawab dari data', 'suggested questions answered from data')), ('0', t('pelanggaran axe serius', 'serious axe violations')), ('≥ 90', t('gate Lighthouse (CI)', 'Lighthouse gate (CI)')), ('0', t('kerentanan npm audit', 'npm audit vulnerabilities'))]
    for i, (v, l) in enumerate(q):
        stat(s, Inches(0.6 + (i % 4) * 3.1), Inches(2.2 + (i // 4) * 2.1), v, l)
    footer(s, n, L)

    # 14 Security & privacy
    s = new()
    title(s, t('Keamanan & privasi', 'Security & privacy'), t('Aman sejak desain', 'Secure by design'))
    bullets(s, Inches(0.6), Inches(2.1), Inches(6.0), Inches(4.8), [
        t('Tanpa akun, tanpa analytics pihak ketiga', 'No accounts, no third-party analytics'),
        t('Riwayat & preferensi hanya di perangkat', 'History & preferences stay on the device'),
        t('Blokir seed phrase, private key, kartu, NIK', 'Blocks seed phrases, private keys, cards, IDs'),
        t('Server tidak menyimpan isi pesan', 'The server does not store message contents'),
    ], 17)
    bullets(s, Inches(6.9), Inches(2.1), Inches(6.0), Inches(4.8), [
        t('CSP ketat, HSTS, CORS allowlist, rate limit', 'Strict CSP, HSTS, CORS allowlist, rate limits'),
        t('Pertahanan prompt injection: data ≠ instruksi', 'Prompt-injection defence: data ≠ instructions'),
        t('gitleaks, npm audit, pemeriksa lisensi', 'gitleaks, npm audit, licence checker'),
        t('Kebijakan privasi mengacu UU PDP 27/2022', 'Privacy policy refers to PDP Law 27/2022'),
    ], 17)
    footer(s, n, L)

    # 15 Deployment
    s = new()
    title(s, t('Deploy', 'Deployment'), t('Dari repositori ke web dan aplikasi', 'From repository to web and apps'))
    table(s, Inches(0.6), Inches(2.0), Inches(12.1), [
        [t('Komponen', 'Component'), t('Status', 'Status')],
        ['GitHub', 'github.com/DSKPutra/finandex_beta · releases data-inputs-v1, law-docs-v1'],
        ['Vercel', t('finandex-beta.vercel.app — produksi (web + API)', 'finandex-beta.vercel.app — production (web + API)')],
        ['Netlify', t('Konfigurasi siap; site finandex-beta menunggu kredit akun', 'Configuration ready; site finandex-beta awaiting account credits')],
        ['GitHub Actions', 'ci · parity · promote · android · ios · deploy-fallback · eval-live'],
        ['Android / iOS', t('Capacitor 8 — APK/AAB & IPA dibangun di CI', 'Capacitor 8 — APK/AAB & IPA built in CI')],
    ], [Inches(3.0), Inches(9.1)], 14, row_h=Inches(0.5))
    footer(s, n, L)

    # 16 Roadmap
    s = new()
    title(s, t('Peta jalan', 'Roadmap'), t('Langkah berikutnya', 'Next steps'))
    flow(s, Inches(2.4), [
        (t('Mode AI', 'AI mode'), t('Aktifkan Claude di produksi; eval live ≥ 95%.', 'Enable Claude in production; live eval ≥ 95%.')),
        (t('Hukum', 'Law'), t('UUD hasil perubahan, KUHPerdata, HIR/RBg, peraturan pelaksana.', 'Amended Constitution, Civil Code, HIR/RBg, implementing rules.')),
        (t('Lisensi', 'Licensing'), t('Konfirmasi lisensi dataset untuk distribusi.', 'Confirm dataset licences for distribution.')),
        (t('Toko aplikasi', 'App stores'), t('Rilis Android & iOS.', 'Release Android & iOS.')),
        (t('Data', 'Data'), t('Pembaruan berkala dengan golden test.', 'Periodic updates with golden tests.')),
    ])
    footer(s, n, L)

    # 17 Closing
    s = new()
    bg(s)
    s.shapes.add_picture(os.path.join(ROOT, 'kit/brand/icons/finandex-1024.png'), Inches(0.8), Inches(0.9), width=Inches(1.0))
    text(s, Inches(0.8), Inches(2.1), Inches(11), Inches(0.9), t('Terima kasih', 'Thank you'), 48, TEXT, HEAD, True)
    text(s, Inches(0.8), Inches(3.1), Inches(11.5), Inches(1.6), ['finandex-beta.vercel.app', 'finandex-beta.vercel.app/whitepaper — PDF ID / EN', 'github.com/DSKPutra/finandex_beta'], 18, GREEN, spacing=1.3)
    text(s, Inches(0.8), Inches(4.7), Inches(11.5), Inches(0.5), t('Finandex dirancang dan dikembangkan oleh Dea Saka Kurnia Putra — A BRIGHTER TOMORROW THROUGH MEANINGFUL WORK.', 'Finandex is designed and developed by Dea Saka Kurnia Putra — A BRIGHTER TOMORROW THROUGH MEANINGFUL WORK.'), 15, TEXT)
    text(s, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.2), t('Data historis untuk informasi dan edukasi — bukan rekomendasi investasi, nasihat keuangan, maupun nasihat hukum. Data tidak real-time.', 'Historical data for information and education — not investment, financial, or legal advice. Data is not real-time.'), 12, MUTED)
    text(s, Inches(0.8), Inches(6.9), Inches(11), Inches(0.4), t('© 2026 Dea Saka Kurnia Putra. Hak cipta dilindungi undang-undang.', '© 2026 Dea Saka Kurnia Putra. All rights reserved.'), 11, MUTED)

    prs.core_properties.title = 'Finandex'
    prs.core_properties.author = 'Dea Saka Kurnia Putra'
    prs.core_properties.subject = t('Presentasi platform Finandex', 'Finandex platform presentation')
    out = os.path.join(ROOT, f"docs/presentation/Finandex-Presentation-{L.upper()}.pptx")
    prs.save(out)
    print(out, n, 'slides')


for L in ('id', 'en'):
    build(L)
